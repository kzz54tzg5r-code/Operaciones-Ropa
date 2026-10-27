"""Regresiones V382: el Excel mostrado en la captura tiene AA/AC/AE.

No contiene registros reales del usuario. Conserva compatibilidad Z/AB/AD.
"""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from daily_sales_headers import detect_monthly_layout, locate_daily_headers
from v370_monthly_store_guard_patch import _missing_store_with_movements


@pytest.mark.parametrize("shift", [0, 1])
def test_identifies_store_dev_and_first_daily_header(shift):
    store=25+shift
    dates={29+shift: date(2026,10,1), 30+shift: date(2026,10,1)}
    labels={
        store:"Tienda",
        store+1:"Ventas Netas Pzs",
        store+2:"Dev Pzs",
        store+3:"Venta Neta en $",
        29+shift:"Ventas Neta Pzs",
        30+shift:"Dev Pzs",
        31+shift:"Venta Neta en $",
    }
    layout=detect_monthly_layout([dates, labels])
    assert layout == {
        "store_col":store, "monthly_dev_col":store+2,
        "first_daily_col":29+shift,
    }
    groups,warnings,_=locate_daily_headers([dates,labels],layout["first_daily_col"])
    assert len(groups)==1 and warnings==[]
    assert groups[0][0]==29+shift


@pytest.mark.parametrize("shift", [0,1])
def test_real_xlsx_streams_daily_dev_without_swapping_columns(tmp_path,shift):
    from web_app import _xlsx_stream_book, _stream_monthly_recovery_fifo
    wb=Workbook()
    ws=wb.active
    ws.title="Octubre 26"
    store_col=26+shift  # one-based: Z or AA
    day_col=30+shift  # one-based: AD or AE
    ws.cell(row=2,column=store_col,value="Tienda")
    ws.cell(row=2,column=store_col+1,value="Ventas Netas Pzs")
    ws.cell(row=2,column=store_col+2,value="Dev Pzs")
    ws.cell(row=2,column=store_col+3,value="Venta Neta en $")
    ws.cell(row=1,column=day_col,value=date(2026,10,1))
    ws.cell(row=2,column=day_col,value="Ventas Neta Pzs")
    ws.cell(row=2,column=day_col+1,value="Dev Pzs")
    ws.cell(row=2,column=day_col+2,value="Venta Neta en $")
    ws.cell(row=3,column=2,value="SKU-0001")
    ws.cell(row=3,column=8,value="NEGRO")
    ws.cell(row=3,column=25,value=129)
    ws.cell(row=3,column=store_col,value="Puebla Sur")
    ws.cell(row=3,column=store_col+1,value=8)
    ws.cell(row=3,column=store_col+2,value=3)
    ws.cell(row=3,column=store_col+3,value=1000)
    ws.cell(row=3,column=day_col,value=8)
    ws.cell(row=3,column=day_col+1,value=3)
    ws.cell(row=3,column=day_col+2,value=1000)
    if shift:
        ws.cell(row=3,column=26,value=499)  # Z tiene costo, NO Tienda
    file=tmp_path/f"monthly_shift_{shift}.xlsx"
    wb.save(file)
    with _xlsx_stream_book(file) as book:
        fifo,weeks,months,errors=_stream_monthly_recovery_fifo(file,["Octubre 26"],book)
        audit=book["recovery_header_audit"]["Octubre 26"]
    assert errors==[],errors
    assert months==["2026-10"]
    assert weeks==["2026-W40"]
    assert audit["store_col"]==25+shift
    assert audit["monthly_dev_col"]==27+shift
    assert audit["first_daily_col"]==29+shift
    assert audit["total_monthly_dev"]==3
    assert audit["total_daily_dev"]==3
    assert len(fifo)==1
    assert fifo[0]["store"]=="Puebla Sur"
    assert fifo[0]["dev_pzs"]==3


def test_missing_store_new_format_is_detected_even_when_z_holds_price():
    row={1:"SKU",25:499,26:"",28:3,30:5,31:3,32:1000}
    assert _missing_store_with_movements(row,store_col=26)
    assert not _missing_store_with_movements({**row,26:"Vallejo"},store_col=26)


def test_old_style_still_uses_z_ab_ad():
    header={25:"Tienda",27:"Dev Pzs",29:"Ventas Neta Pzs",30:"Dev Pzs",31:"Venta Neta en $"}
    layout=detect_monthly_layout([{},header])
    assert layout["store_col"]==25
    assert layout["monthly_dev_col"]==27
    assert layout["first_daily_col"]==29
