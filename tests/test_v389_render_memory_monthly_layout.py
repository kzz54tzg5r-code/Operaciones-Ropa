"""Regression: FIFO compacto alinea DEV con fecha y tienda para 2 layouts."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
import pytest


def _book(path, shifted):
    wb=Workbook()
    ws=wb.active
    ws.title="Octubre 26"
    store="AA" if shifted else "Z"
    monthly="AC" if shifted else "AB"
    first=("AE","AF","AG","AH","AI","AJ") if shifted else ("AD","AE","AF","AG","AH","AI")
    ws[f"{store}2"]="Tienda"
    ws[f"{monthly}2"]="Dev Pzs"
    for day,triplet in ((24,first[:3]),(25,first[3:])):
        for col in triplet:
            ws[f"{col}1"]=date(2026,10,day)
        for col,metric in zip(triplet,("Ventas Neta Pzs","Dev Pzs","Venta Neta en $")):
            ws[f"{col}2"]=metric
    ws["B3"]="SKU-1"
    ws["Y3"]=125
    ws[f"{store}3"]="PUEBLA SUR"
    ws[f"{monthly}3"]=4
    for col,value in zip(first,(30,2,500,90,2,1500)):
        ws[f"{col}3"]=value
    wb.save(path)


@pytest.mark.parametrize("shifted", [False, True])
def test_compact_monthly_fifo_not_advance_sales_as_previous_day_dev(tmp_path,shifted):
    import web_app as m
    import render_memory_patch
    render_memory_patch.install(m)
    xlsx=tmp_path/"octubre.xlsx"
    _book(xlsx,shifted)
    fifo,weeks,months,errors=m._stream_monthly_recovery_fifo(xlsx,["Octubre 26"])
    assert not errors,errors
    bydate={r["date"]:r for r in fifo}
    assert bydate["2026-10-24"]["dev_pzs"]==2
    assert bydate["2026-10-25"]["dev_pzs"]==2
    assert sum(r["dev_pzs"] for r in fifo)==4
    assert all(r["store"]=="Puebla Sur" for r in fifo)
    assert "2026-10" in months
    assert "2026-W43" in weeks


def test_render_layout_parser_is_header_driven():
    text=(Path(__file__).resolve().parents[1]/"render_memory_patch.py").read_text(encoding="utf-8")
    assert "detect_monthly_layout(header_rows)" in text
    assert "locate_daily_headers(" in text
    assert 'cell(values, store_col' in text
    assert "for index in range(29, max_column" not in text
