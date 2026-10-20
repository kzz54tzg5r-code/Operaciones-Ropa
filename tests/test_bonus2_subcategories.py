from pathlib import Path
import pandas as pd
import pytest
from v340_bonus2_subcategories import coverage, subcategory, compare_rows, inventory, store_metric, select_targets


def article(ident='1', category='Chamarra', stock=200, demand=1, investment=200):
    return dict(store='Olivar',id_art=ident,model='M'+ident,catalog=category,
                existence=stock,sug7=demand,investment=investment,ddi=coverage(stock,demand))


def test_categories_and_demand_units():
    assert [subcategory(x) for x in ('CHAMARRAS','Sudadera','jeans','Abrigador')] == ['Chamarra','Sudadera','Jeans',None]
    assert coverage(100,2)==50  # no extra factor of seven
    assert coverage(0,2)==0
    assert coverage(100,0) is None
    assert coverage(100,None) is None


def test_fixed_demand_does_not_reward_change_in_sug7_or_double_count_deliveries():
    initial=article(stock=200,demand=1)
    final=article(stock=150,demand=3)
    row=compare_rows([initial],[final])[0]
    assert row['initial_ddi']==200
    assert row['current_ddi']==50
    assert row['fixed_ddi']==150
    assert row['ddi_reduction_pct']==25
    assert row['current_existence']==150
    assert row['deliveries'] is None
    assert inventory([row])['current_existence']==150


def test_missing_article_or_second_cut_never_becomes_zero_inventory():
    for current in (None,[],[article('2')]):
        row=compare_rows([article()],current)[0]
        assert row['current_existence'] is None
        assert row['fixed_ddi'] is None
        assert not row['ddi_reduced']
        assert row['ddi_reduction_pct'] is None
    zero=compare_rows([article()],[article(stock=0)])[0]
    assert zero['fixed_ddi']==0
    assert zero['ddi_reduced']


def test_same_target_selection_per_subcategory_and_investment_informative():
    xs=[article('1',investment=300),article('2',investment=100),article('3','Jeans',investment=10000)]
    assert [x['id_art'] for x in select_targets(xs)]==['1']
    row=compare_rows([xs[0]],[article('1',stock=150,investment=999999)])[0]
    source=dict(store='Olivar',sales_pct=100,attendance_pct=100,attendance_provisional=False)
    result=store_metric(source,[row])
    assert result['ddi_pct']==100
    assert result['score']==100
    assert result['score_ready']
    assert store_metric(dict(source,attendance_provisional=True),[row])['score_ready'] is False


def test_missing_one_subcategory_cannot_make_partial_ddi_complete():
    rows=compare_rows([article('1'),article('2','Jeans')],[article('1',stock=100)])
    result=store_metric(dict(store='Olivar',sales_pct=0,attendance_pct=0,attendance_provisional=False),rows)
    assert result['ddi_pct'] is None
    assert result['score']==0
    assert not result['score_ready']
    assert inventory(rows)['current_existence'] is None


def test_source_uses_subcategory_not_commercial_catalog(tmp_path):
    from v340_bonus2_subcategories import read_snapshot
    path=tmp_path/'capacity.xlsx'
    pd.DataFrame({'TIENDA':['Olivar']*3,'ID_ART':[10001,10002,10003],'MODELO':['A','B','C'],
                  'SUBCATEGORIA':['Chamarra','Sudadera','Jeans'],'CATALOGO':['OTRO']*3,
                  'TIPO CATALOGO MAX VIG':['VIGENTE']*3,'SUG 7':[1,0,2],
                  'EXISTENCIA TOTAL':[200,50,300],'PRECIO MAYOREO':[100]*3}).to_excel(path,index=False)
    rows=read_snapshot(path)
    assert {r['catalog'] for r in rows}=={'Chamarra','Sudadera','Jeans'}
    assert next(r for r in rows if r['catalog']=='Sudadera')['ddi'] is None


def test_endpoint_scope_and_exports_share_all_rows(tmp_path, monkeypatch):
    import json
    from datetime import date
    from types import SimpleNamespace
    from fastapi import FastAPI, HTTPException
    from v340_bonus2_subcategories import install
    import v340_bonus2_subcategories as module
    from v322_bonus_model_exports import install as exports
    from io import BytesIO
    from openpyxl import load_workbook
    app=FastAPI()
    actor={'role':'tienda','store':'Olivar'}
    @app.get('/api/operation/sales-bonus-v293')
    def base(request=None,month='',store=''):
        return dict(month='2026-10',months=[{'value':'2026-10','label':'Octubre 2026'}],stores=['Olivar'],selected_store='Olivar',
                    source_dates={'opening_capacity':'2026-10-01','current_capacity':'2026-10-15','current_capacity_id':'final'},
                    ranking=[dict(store='Olivar',sales_pct=100,sales_actual=100,sales_goal=100,attendance_pct=100,attendance_provisional=False)])
    first=tmp_path/'initial.xlsx';first.write_bytes(b'initial')
    last=tmp_path/'final.xlsx';last.write_bytes(b'final')
    entries=[{'id':'initial','path':str(first)},{'id':'final','path':str(last)}]
    (tmp_path/'manifest.json').write_text(json.dumps({'capacities':entries}))
    monkeypatch.setattr(module,'read_snapshot',lambda p: [article(str(10000+i),investment=300 if i<35 else 10,
                         stock=200 if p==first else 100) for i in range(40)]+[dict(article('99999'),store='Otra tienda')])
    m=SimpleNamespace(app=app,DATA_ROOT=tmp_path,COMMERCIAL_DATA_ROOT=tmp_path,require_user=lambda r:actor,
                      resolve_entry_path=lambda e:Path(e['path']),_sales_bonus_opening_capacity_v293=lambda month:(date(2026,10,1),None,entries[0]))
    install(m)
    import time
    for _ in range(100):
        payload=m._v340_bonus2_payload(None,'2026-10','Compañía')
        if not payload['processing']['capacity']: break
        time.sleep(.01)
    assert payload['selected_store']=='Olivar'
    assert payload['stores']==['Olivar']
    assert len(payload['details'])==35
    assert all(r['store']=='Olivar' for r in payload['details'])
    assert payload['bonus2_rows'][0]['score']==100
    with pytest.raises(HTTPException): m._v340_bonus2_payload(None,'2026-10','Otra tienda')
    exports(m)
    endpoint=next(r.endpoint for r in app.routes if getattr(r,'path','')=='/api/bonuses/models-export')
    result=endpoint(None,'2026-10','Olivar',2,'Chamarra','xlsx',False)
    sheet=load_workbook(BytesIO(result.body)).active
    assert sheet.max_row==39  # all 35 models, not the first 30 shown on screen
    assert sheet['K5'].value==100
    assert sheet['G5'].value==100
    assert 'Sin fuente' in sheet['F5'].value
    pdf=endpoint(None,'2026-10','Olivar',2,'Chamarra','pdf',False)
    assert pdf.body.startswith(b'%PDF-')
