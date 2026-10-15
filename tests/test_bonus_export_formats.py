"""Contrato real de Excel/PDF para Bonos -> Modelos a Detalle."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook

from v322_bonus_model_exports import install


def app_with_sample():
    mock=SimpleNamespace(app=FastAPI(), require_user=lambda request: {"role":"admin"})

    @mock.app.get("/api/bonuses-v313")
    def get_models(request=None,month="",store="",bonus=1):
        return {
            "selected_store":store,
            "month":month,
            "details":[
                {
                    "id_art":"108180","model":"9112","catalog":"Abrigador",
                    "initial_ddi":280,"current_ddi":220,"ddi_reduction_pct":21.4,
                    "initial_existence":243,"current_existence":221,
                    "initial_investment":60507,"current_investment":54001,
                    "investment_reduction_pct":10.7,
                },
                {
                    "id_art":"1203206","model":"880","catalog":"Licencias",
                    "initial_ddi":109,"current_ddi":105,"ddi_reduction_pct":3.67,
                    "initial_existence":150,"current_existence":145,
                    "initial_investment":29000,"current_investment":28000,
                    "investment_reduction_pct":3.44,
                },
            ],
        }
    install(mock)
    return next(r.endpoint for r in mock.app.routes if getattr(r,"path",None)=="/api/bonuses/models-export")


def test_excel_is_real_xlsx_and_respects_catalog():
    export=app_with_sample()
    response=export(request=None,month="2026-10",store="Olivar",bonus=1,catalog="Abrigador",format="xlsx",photos=False)
    assert response.status_code==200
    assert response.body[:2]==b"PK"
    assert response.headers["content-disposition"].endswith('.xlsx"')
    sheet=load_workbook(BytesIO(response.body))["Modelos a Detalle"]
    assert sheet["B5"].value=="108180"
    assert sheet["D5"].value=="Abrigador"
    assert sheet["B6"].value is None


def test_pdf_is_binary_pdf_with_download_header():
    export=app_with_sample()
    response=export(request=None,month="2026-10",store="Olivar",bonus=1,catalog="Licencias",format="pdf",photos=False)
    assert response.status_code==200
    assert response.body.startswith(b"%PDF-")
    assert response.headers["content-disposition"].endswith('.pdf"')


def test_no_models_returns_clear_http_error_without_fake_file():
    export=app_with_sample()
    with pytest.raises(HTTPException) as captured:
        export(request=None,month="2026-10",store="Olivar",bonus=1,catalog="Básicos",format="xlsx",photos=False)
    assert captured.value.status_code==404
    assert "No hay modelos" in captured.value.detail


def test_excel_and_pdf_embed_the_same_approved_web_photo():
    """Photos visible through the authenticated web endpoint also reach exports."""
    from PIL import Image
    from pypdf import PdfReader
    sample=BytesIO()
    Image.new("RGB",(80,110),(31,92,157)).save(sample,format="JPEG")
    raw=sample.getvalue()
    app=FastAPI()
    mock=SimpleNamespace(
        app=app,
        require_user=lambda request: {"role":"admin"},
        _bonus_photo_fetch_v314=lambda ident: (raw,"image/jpeg") if ident=="108180" else None
    )

    @app.get("/api/bonuses-v313")
    def endpoint(request=None,month="",store="",bonus=1):
        return {"selected_store":store,"month":month,
                "details":[{"id_art":"108180","model":"9112","catalog":"Abrigador",
                            "initial_ddi":280,"current_ddi":220,"ddi_reduction_pct":21.4,
                            "initial_existence":243,"current_existence":221,
                            "initial_investment":60507,"current_investment":54001,
                            "investment_reduction_pct":-10.7}]}
    install(mock)
    export=next(r.endpoint for r in app.routes if getattr(r,"path",None)=="/api/bonuses/models-export")
    excel=export(request=None,store="Ixtapaluca",month="2026-10",photos=True,format="xlsx")
    assert excel.headers["x-photos-included"] == "1"
    sheet=load_workbook(BytesIO(excel.body))["Modelos a Detalle"]
    assert len(sheet._images)==1
    assert sheet["A5"].value is None
    pdf=export(request=None,store="Ixtapaluca",month="2026-10",photos=True,format="pdf")
    assert pdf.headers["x-photos-included"] == "1"
    page=PdfReader(BytesIO(pdf.body)).pages[0]
    text=page.extract_text()
    assert "-10.7%" in text  # It must not display as $-10.7.
    assert "60,507.00" in text
    assert page.mediabox.width > page.mediabox.height
    assert len(page.images) >= 1


def test_pdf_uses_entire_available_page_width_for_large_tables(monkeypatch):
    """At least the title and all twelve columns must fit on the A3 landscape page."""
    from pypdf import PdfReader
    from v322_bonus_model_exports import FIELDS
    app=SimpleNamespace(app=FastAPI(),require_user=lambda request:{"role":"admin"})
    @app.app.get("/api/bonuses-v313")
    def get_models(request=None,month="",store="",bonus=1):
        return {"selected_store":store,"month":month,"details":[{
            "id_art":str(1037708+i), "model":str(4620+i),"catalog":"Abrigador",
            "initial_ddi":329,"current_ddi":343,"ddi_reduction_pct":-4.3,
            "initial_existence":47,"current_existence":49,
            "initial_investment":39433,"current_investment":41111,"investment_reduction_pct":-4.3
        } for i in range(30)]}
    install(app)
    export=next(r.endpoint for r in app.app.routes if getattr(r,"path",None)=="/api/bonuses/models-export")
    resp=export(request=None,store="Ixtapaluca",month="2026-10",catalog="Abrigador",format="pdf",photos=False)
    reader=PdfReader(BytesIO(resp.body))
    assert reader.pages[0].mediabox.width == pytest.approx(1190.55,abs=1)
    txt=reader.pages[0].extract_text()
    assert "Inversión actual" in txt
    assert "Reducción inversión" in txt
    assert "-4.3%" in txt



def test_real_photo_patch_installation_is_wired_into_exports(monkeypatch):
    """Integration regression: the shared V314 provider must actually be exposed
    to the V322 export installer, not commented out via literal backslash-n.
    """
    import v314_bonus_product_photos_patch as module
    from PIL import Image
    from pypdf import PdfReader

    sample=BytesIO()
    Image.new("RGB",(80,100),(30,90,140)).save(sample,format="JPEG")
    raw=sample.getvalue()

    class FakeRemoteImage:
        headers={"Content-Type":"image/jpeg"}
        def __enter__(self): return self
        def __exit__(self,*args): return False
        def read(self,length): return raw

    calls=[]
    def fake_request(request,timeout=5):
        calls.append(request.full_url)
        return FakeRemoteImage()
    monkeypatch.setattr(module.urllib.request, "urlopen", fake_request)

    m=SimpleNamespace(app=FastAPI(),require_user=lambda request:{"role":"admin"})
    @m.app.get("/api/bonuses-v313")
    def bonus_data(request=None,month="",store="",bonus=1):
        return {
            "selected_store":store,"month":month,
            "details":[{
                "id_art":"1888555","model":"C100","catalog":"Abrigador",
                "initial_ddi":120,"current_ddi":100,"ddi_reduction_pct":16.67,
                "initial_existence":22,"current_existence":19,
                "initial_investment":2500,"current_investment":2200,
                "investment_reduction_pct":12.0
            }]
        }
    module.install(m)
    assert callable(m._bonus_photo_fetch_v314)
    install(m)
    export=next(rt.endpoint for rt in m.app.routes if getattr(rt,"path",None)=="/api/bonuses/models-export")
    response_xlsx=export(request=None,month="2026-10",store="Ixtapaluca",catalog="Abrigador",format="xlsx")
    assert response_xlsx.headers["x-photos-included"] == "1"
    assert response_xlsx.headers["x-photos-requested"] == "1"
    sheet=load_workbook(BytesIO(response_xlsx.body)).active
    assert len(sheet._images)==1
    response_pdf=export(request=None,month="2026-10",store="Ixtapaluca",catalog="Abrigador",format="pdf")
    assert response_pdf.headers["x-photos-included"] == "1"
    assert any(page.images for page in PdfReader(BytesIO(response_pdf.body)).pages)
    assert calls and all(url.endswith("1888555.jpg") for url in calls)


def test_when_photo_server_is_unavailable_photo_exports_fail_explicitly():
    m=SimpleNamespace(
        app=FastAPI(),
        require_user=lambda request:{"role":"admin"},
        _bonus_photo_fetch_v314=lambda ident:None,
    )
    @m.app.get("/api/bonuses-v313")
    def bonus_data(request=None,month="",store="",bonus=1):
        return {"selected_store":store,"month":month,"details":[{"id_art":"1888555","model":"C100","catalog":"Abrigador"}]}
    install(m)
    export=next(rt.endpoint for rt in m.app.routes if getattr(rt,"path",None)=="/api/bonuses/models-export")
    with pytest.raises(HTTPException) as exc:
        export(request=None,store="Ixtapaluca",catalog="Abrigador",format="xlsx",photos=True)
    assert exc.value.status_code == 503
    assert "ninguna fotografía" in exc.value.detail
    # An explicit photos=False request still produces a readable workbook.
    resp=export(request=None,store="Ixtapaluca",catalog="Abrigador",format="xlsx",photos=False)
    assert resp.status_code == 200


def test_export_filename_includes_store_catalog_and_mexico_date():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    expected=datetime.now(ZoneInfo("America/Mexico_City")).strftime("%d.%m.%y")
    export=app_with_sample()
    for format in ("pdf","xlsx"):
        response=export(request=None,month="2026-10",store="Ixtapaluca",bonus=1,
                        catalog="Abrigador",format=format,photos=False)
        assert response.headers["content-disposition"]==(
            'attachment; filename="Modelos Ixtapaluca Abrigador '+expected+'.'+format+'"'
        )


def test_dedicated_export_photo_source_handles_more_than_one_model_and_both_file_types():
    from PIL import Image
    from pypdf import PdfReader
    sample=BytesIO()
    Image.new("RGB",(80,105),(30,60,90)).save(sample,format="JPEG")
    img=sample.getvalue()
    sample2=BytesIO()
    Image.new("RGB",(80,105),(100,150,25)).save(sample2,format="JPEG")
    img2=sample2.getvalue()
    m=SimpleNamespace(
        app=FastAPI(),require_user=lambda req:{"role":"admin"},
        _bonus_photo_export_v314=lambda product: img if product=="108180" else (img2 if product=="1203206" else None),
    )
    @m.app.get("/api/bonuses-v313")
    def sample(request=None,month="",store="",bonus=1):
        return {"selected_store":store,"month":month,"details":[
            {"id_art":"108180","model":"9112","catalog":"Abrigador"},
            {"id_art":"1203206","model":"880","catalog":"Abrigador"},
        ]}
    install(m)
    export=next(r.endpoint for r in m.app.routes if getattr(r,"path",None)=="/api/bonuses/models-export")
    x=export(request=None,month="2026-10",store="Ixtapaluca",format="xlsx",photos=True)
    assert x.headers["x-photos-included"]=="2"
    assert len(load_workbook(BytesIO(x.body)).active._images)==2
    p=export(request=None,month="2026-10",store="Ixtapaluca",format="pdf",photos=True)
    assert p.headers["x-photos-included"]=="2"
    assert sum(len(page.images) for page in PdfReader(BytesIO(p.body)).pages)>=2
