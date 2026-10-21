"""Concurrency, authorization, freshness and output contracts for runtime work."""
import asyncio
from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import sqlite3
import threading
import time
from types import SimpleNamespace

from fastapi import FastAPI, HTTPException, Request
from fastapi.testclient import TestClient
import httpx
import pytest

from runtime_performance import ReportCache, SerialJobQueue, data_revision, install


def make_app(tmp_path, endpoint=None):
    app = FastAPI()
    counts = {"calls": 0, "auth": 0}

    def actor(request):
        counts["auth"] += 1
        name = request.headers.get("x-user")
        if not name:
            raise HTTPException(401, "Sesión requerida")
        return {"id": name, "username": name, "role": request.headers.get("x-role", "admin"),
                "store": request.headers.get("x-assigned-store", "")}

    @app.get("/api/operations")
    def report(request: Request, store: str = "Compañía", period: str = "2026-10"):
        counts["calls"] += 1
        user = actor(request)
        if endpoint:
            return endpoint(request, store, period)
        selected = user["store"] if user["role"] == "tienda" else store
        source = tmp_path / "source.json"
        data = json.loads(source.read_text()) if source.exists() else {"pieces": 125}
        return {"store": selected, "period": period, "pieces": data["pieces"],
                "percent": 42.6, "absent": None, "zero": 0}

    @app.get("/api/bonuses/models-export")
    def download(request: Request):
        actor(request)
        from fastapi.responses import Response
        return Response(b"%PDF-" + b"x" * 2000, media_type="application/pdf")

    module = SimpleNamespace(app=app, DATA_ROOT=tmp_path, Request=Request,
                             require_user=actor, require_real_superadmin=actor,
                             _UPLOAD_QUEUE=SerialJobQueue())
    original = next(r for r in app.routes if r.path == "/api/operations").endpoint
    install(module)
    return app, module, counts, original


def test_reports_keep_numbers_and_scope_and_check_auth_on_every_hit(tmp_path):
    app, module, counts, original = make_app(tmp_path)
    with TestClient(app) as client:
        url = "/api/operations?store=Vallejo&period=2026-10"
        first = client.get(url, headers={"x-user": "a"})
        second = client.get(url, headers={"x-user": "a"})
        assert first.content == second.content
        assert first.json()["percent"] == 42.6
        assert first.json()["absent"] is None and first.json()["zero"] == 0
        assert second.headers["x-report-cache"] == "HIT"
        assert counts["calls"] == 1
        assert client.get(url).status_code == 401
        for headers in ({"x-user": "b"}, {"x-user": "a", "x-role": "tienda", "x-assigned-store": "Olivar"}):
            result = client.get(url, headers=headers)
            assert result.headers["x-report-cache"] == "MISS"
        assert result.json()["store"] == "Olivar"
        assert client.get(url + "&period=2026-11", headers={"x-user": "a"}).headers["x-report-cache"] == "MISS"
        bypass = client.get(url, headers={"x-user": "a", "cache-control": "no-cache"})
        assert bypass.headers["x-report-cache"] == "BYPASS"
        assert bypass.content == first.content
        forward = client.get("/api/operations?period=2026-10&period=2026-11", headers={"x-user": "a"})
        reverse = client.get("/api/operations?period=2026-11&period=2026-10", headers={"x-user": "a"})
        assert forward.json()["period"] == "2026-11"
        assert reverse.json()["period"] == "2026-10"
        assert next(r for r in app.routes if r.path == "/api/operations").endpoint is original


def test_cache_invalidates_after_new_data_and_wal_commit(tmp_path):
    app, module, counts, _ = make_app(tmp_path)
    db = sqlite3.connect(tmp_path / "users.sqlite3")
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("CREATE TABLE settings(value)")
    db.commit()
    with TestClient(app) as client:
        headers = {"x-user": "a"}
        assert client.get("/api/operations", headers=headers).json()["pieces"] == 125
        assert client.get("/api/operations", headers=headers).headers["x-report-cache"] == "HIT"
        (tmp_path / "source.json").write_text('{"pieces":300}')
        updated = client.get("/api/operations", headers=headers)
        assert updated.json()["pieces"] == 300
        assert updated.headers["x-report-cache"] == "MISS"
        assert client.get("/api/operations", headers=headers).headers["x-report-cache"] == "HIT"
        before = data_revision(tmp_path)
        db.execute("INSERT INTO settings VALUES(47)")
        db.commit()
        assert data_revision(tmp_path) != before
        assert client.get("/api/operations", headers=headers).headers["x-report-cache"] == "MISS"
    db.close()


def test_parallel_identical_requests_compute_once_and_keep_threadpool_free(tmp_path):
    def slow(*_):
        time.sleep(.08)
        return {"pieces": 784, "attendance": 20}
    app, module, counts, _ = make_app(tmp_path, slow)

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            responses = await asyncio.gather(*[
                client.get("/api/operations?store=Olivar", headers={"x-user": "admin"})
                for _ in range(30)
            ])
        assert all(r.status_code == 200 and r.json() == {"pieces": 784, "attendance": 20} for r in responses)
        assert counts["calls"] == 1
        assert module._REPORT_PERFORMANCE_CACHE.stats()["shared"] > 0
    asyncio.run(exercise())


def test_distinct_expensive_reports_have_bounded_concurrency(tmp_path):
    lock = threading.Lock()
    active = maximum = 0
    def slow(*_):
        nonlocal active, maximum
        with lock:
            active += 1
            maximum = max(maximum, active)
        time.sleep(.03)
        with lock:
            active -= 1
        return {"ok": True}
    app, _, _, _ = make_app(tmp_path, slow)
    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            result = await asyncio.gather(*[client.get(f"/api/operations?period={i}", headers={"x-user": "a"}) for i in range(10)])
            assert all(r.status_code == 200 for r in result)
    asyncio.run(exercise())
    assert maximum == 2


def test_async_report_authorities_are_cached_too(tmp_path):
    app, module, _, _ = make_app(tmp_path)
    calls = []
    app.router.routes = [r for r in app.routes if r.path != "/api/operations"]

    @app.get("/api/operations")
    async def authority(request: Request):
        calls.append(1)
        return {"pieces": 784, "percentage": 42.6}

    # Simulates the final async authority installed by the production patches.
    module._RUNTIME_PERFORMANCE_INSTALLED = False
    install(module)
    with TestClient(app) as client:
        first = client.get("/api/operations", headers={"x-user": "a"})
        second = client.get("/api/operations", headers={"x-user": "a"})
        assert first.json() == second.json() == {"pieces": 784, "percentage": 42.6}
        assert second.headers["x-report-cache"] == "HIT"
        assert len(calls) == 1


def test_slow_operational_computation_keeps_the_event_loop_available():
    from v153_project_scope_backend_authoritative_patch import install as authority_install
    app = FastAPI()
    started, release = threading.Event(), threading.Event()

    def slow(*_, **__):
        started.set()
        assert release.wait(3)
        return {"pieces": 784}

    module = SimpleNamespace(app=app, project_store_names=lambda *_: ["Olivar"], operations=slow)
    authority_install(module)

    @app.get("/health")
    async def health():
        return {"ok": True}

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            computation = asyncio.create_task(client.get("/api/operations?project_only=true"))
            try:
                assert await asyncio.to_thread(started.wait, 2)
                assert (await asyncio.wait_for(client.get("/health"), .5)).json() == {"ok": True}
                assert not computation.done(), "Health must finish while the report is still computing"
            finally:
                release.set()
            assert (await computation).json() == {"pieces": 784}
    asyncio.run(exercise())


def test_errors_and_loading_are_not_cached_and_binary_exports_stay_intact(tmp_path):
    state = {"mode": "error"}
    def result(*_):
        if state["mode"] == "error":
            raise HTTPException(503, "Iniciando")
        if state["mode"] == "processing":
            return {"processing": {"capacity": True}, "pieces": None}
        return {"pieces": 0}
    app, _, counts, _ = make_app(tmp_path, result)
    with TestClient(app) as client:
        headers = {"x-user": "a"}
        assert client.get("/api/operations", headers=headers).status_code == 503
        state["mode"] = "processing"
        assert client.get("/api/operations", headers=headers).json()["pieces"] is None
        state["mode"] = "ready"
        assert client.get("/api/operations", headers=headers).json()["pieces"] == 0
        assert counts["calls"] == 3
        binary = client.get("/api/bonuses/models-export", headers={**headers, "accept-encoding": "gzip"})
        assert binary.content == b"%PDF-" + b"x" * 2000
        assert "content-encoding" not in binary.headers


@pytest.mark.parametrize("format", ["pdf", "xlsx"])
def test_operational_exports_keep_their_full_output_without_blocking_health(format):
    from io import BytesIO
    from openpyxl import load_workbook
    import pandas as pd
    from v153_project_scope_backend_authoritative_patch import install as authority_install
    app = FastAPI()
    started, release = threading.Event(), threading.Event()
    data = {"metrics": {}, "stores": [{"store": "Olivar"}], "recovery_by_store": []}

    def render_pdf(*_, **__):
        started.set()
        assert release.wait(3)
        return b"%PDF-full-output"

    def rows(*_):
        started.set()
        assert release.wait(3)
        return [["Total", 784]], [{"store": "Olivar", "pieces": i} for i in range(500)], [], []

    module = SimpleNamespace(app=app, project_store_names=lambda *_: ["Olivar"],
                             operations=lambda *_, **__: dict(data),
                             require_user=lambda *_: {"role": "admin"},
                             effective_store=lambda actor, selected: selected,
                             _build_operations_pdf=render_pdf, _report_export_payload=rows, pd=pd)
    authority_install(module)
    @app.get("/health")
    async def health():
        return {"ok": True}

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            task = asyncio.create_task(client.get(f"/api/export/operations/project-scope-v151?store=Olivar&format={format}"))
            try:
                assert await asyncio.to_thread(started.wait, 2)
                assert (await asyncio.wait_for(client.get("/health"), .5)).json()["ok"]
                assert not task.done()
            finally:
                release.set()
            response = await task
            assert response.status_code == 200
            if format == "pdf":
                assert response.content == b"%PDF-full-output"
            else:
                sheet = load_workbook(BytesIO(response.content))["Detalle Proyecto"]
                assert sheet.max_row == 501
                assert sheet["B501"].value == 499
    asyncio.run(exercise())


def test_large_text_is_compressed_without_changing_json(tmp_path):
    payload = {"rows": [{"store": "Iztapalapa", "pieces": 784}] * 1000}
    app, _, _, _ = make_app(tmp_path, lambda *_: payload)
    with TestClient(app) as client:
        result = client.get("/api/operations", headers={"x-user": "a", "accept-encoding": "gzip"})
        assert result.json() == payload
        assert result.headers["content-encoding"] == "gzip"
        assert int(result.headers["content-length"]) < len(result.content) / 5


def test_cache_is_limited_by_bytes_entries_and_time():
    now = [0]
    cache = ReportCache(max_bytes=12, max_entries=2, max_item_bytes=8, ttl=20, clock=lambda: now[0])
    for key in "abc":
        assert cache.acquire(key)[0] == "MISS"
        cache.finish(key, b"123456")
    assert cache.stats()["bytes"] == 12
    assert cache.stats()["entries"] == 2
    assert cache.acquire("a")[0] == "MISS"
    cache.finish("a", b"x" * 9)
    assert cache.stats()["entries"] == 2
    now[0] = 21
    assert cache.acquire("c")[0] == "MISS"
    cache.finish("c", b"x")
    assert cache.stats()["bytes"] == 1


def test_upload_queue_preserves_order_deduplicates_and_continues_after_error():
    release, started = threading.Event(), threading.Event()
    calls, failures = [], []
    queue = SerialJobQueue(on_error=lambda ident, exc: failures.append(ident))
    def first():
        started.set()
        assert release.wait(3)
        calls.append(1)
    def broken():
        calls.append(2)
        raise ValueError("invalid Excel")
    queue.submit("1", first, ())
    assert started.wait(2)
    assert queue.submit("1", first, ()) is False
    queue.submit("2", broken, ())
    queue.submit("3", lambda: calls.append(3), ())
    assert queue.stats() == {"running": 1, "queued": 2}
    release.set()
    queue._queue.join()
    assert calls == [1, 2, 3] and failures == ["2"]
    assert queue.stats() == {"running": 0, "queued": 0}


def test_database_context_closes_commits_and_rolls_back(monkeypatch, tmp_path):
    import web_app as m
    monkeypatch.setattr(m, "DB", tmp_path / "context.sqlite3")
    with m.db() as connection:
        connection.execute("CREATE TABLE sample(value)")
        connection.execute("INSERT INTO sample VALUES(1)")
    with pytest.raises(sqlite3.ProgrammingError):
        connection.execute("SELECT 1")
    with pytest.raises(ValueError):
        with m.db() as connection:
            connection.execute("INSERT INTO sample VALUES(2)")
            raise ValueError("rollback")
    with m.db() as connection:
        assert connection.execute("SELECT value FROM sample").fetchall()[0][0] == 1
        assert connection.execute("SELECT COUNT(*) FROM sample").fetchone()[0] == 1


def test_sales_pdf_job_returns_before_parsing_and_keeps_full_result(monkeypatch, tmp_path):
    import web_app as m
    database = tmp_path / "jobs.sqlite3"
    monkeypatch.setattr(m, "DB", database)
    monkeypatch.setattr(m, "STAGING_DIR", tmp_path)
    m.init_db()
    m.ensure_upload_job_columns()
    release, started = threading.Event(), threading.Event()
    queue = SerialJobQueue()
    monkeypatch.setattr(m, "_UPLOAD_QUEUE", queue)
    monkeypatch.setattr(m, "require_user", lambda *_: {"username": "admin", "role": "admin"})
    captured = []
    def parse(sources, year, month):
        captured.extend(list(sources))
        captured.append((year, month))
        started.set()
        assert release.wait(3)
        return {"results": [{"file": "ventas.pdf", "ok": True, "total_sales": 125000.25, "rows": 784}]}
    monkeypatch.setattr(m, "_process_sales_pdf_sources", parse)
    with TestClient(m.app) as client:
        response = client.post("/api/upload/jobs/sales-pdfs", data={"year": 2026, "month": 10},
                               files={"files": ("ventas.pdf", b"%PDF-fixture", "application/pdf")})
        assert response.status_code == 202
        job = response.json()
        assert started.wait(2) and m._get_upload_job(job["id"])["status"] == "processing"
        assert client.get("/health").status_code == 200
        release.set()
        queue._queue.join()
        finished = m._get_upload_job(job["id"])
        assert finished["status"] == "complete"
        assert finished["result"]["results"][0]["total_sales"] == 125000.25
        assert captured == [("ventas.pdf", b"%PDF-fixture"), (2026, 10)]
        assert not list(tmp_path.glob("sales_pdf_job_*"))


def test_pdf_job_is_resumed_with_the_same_persisted_files(monkeypatch, tmp_path):
    import web_app as m
    monkeypatch.setattr(m, "DB", tmp_path / "jobs.sqlite3")
    m.init_db()
    m.ensure_upload_job_columns()
    pdf = tmp_path / "retained.pdf"
    pdf.write_bytes(b"%PDF-original")
    source = tmp_path / "job.json"
    source.write_text(json.dumps({"year": 2026, "month": 10, "files": [{"path": str(pdf), "name": "original.pdf"}]}))
    job = m._create_upload_job("sales_pdfs", "original.pdf", source, "admin")
    submitted = []
    monkeypatch.setattr(m, "_submit_upload_job", lambda *args: submitted.append(args) or True)
    monkeypatch.setattr(m, "load_manifest", lambda: {})
    assert m._resume_pending_upload_jobs() == 1
    assert submitted == [(job["id"], m._run_sales_pdf_job, (job["id"], source))]
    assert pdf.read_bytes() == b"%PDF-original"
