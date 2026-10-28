"""V387: evitar reprocesos repetidos y asegurar corte y respaldo."""
import json
from pathlib import Path
from types import SimpleNamespace
from datetime import date
from unittest.mock import patch

from openpyxl import Workbook

from v387_one_time_dev_repair import _date_covered_in_original, install, MARKER_NAME


def fixture_book(path, with_day=True):
    wb=Workbook()
    ws=wb.active
    ws.title="Octubre 26"
    ws["AA2"]="Tienda"
    ws["AC2"]="Dev Pzs"
    ws["AE2"]="Ventas Neta Pzs"
    ws["AF2"]="Dev Pzs"
    ws["AG2"]="Venta Neta en $"
    ws["AE1"]=date(2026,10,25 if with_day else 12)
    wb.save(path)


def test_raw_workbook_has_to_contain_specific_date(tmp_path):
    from web_app import _xlsx_stream_book, _xlsx_monthly_rows, _monthly_sheet_name
    excel=tmp_path/"cambios_muertos_actual.xlsx"
    m=SimpleNamespace(_xlsx_stream_book=_xlsx_stream_book,
                      _xlsx_monthly_rows=_xlsx_monthly_rows,
                      _monthly_sheet_name=_monthly_sheet_name)
    fixture_book(excel,with_day=True)
    assert _date_covered_in_original(m,excel)[0]
    fixture_book(excel,with_day=False)
    assert not _date_covered_in_original(m,excel)[0]


def test_one_attempt_per_raw_fingerprint(tmp_path):
    from web_app import _xlsx_stream_book, _xlsx_monthly_rows, _monthly_sheet_name
    excel=tmp_path/"cambios_muertos_actual.xlsx"
    fixture_book(excel)
    ops=tmp_path/"operaciones.json"
    ops.write_text("{}",encoding="utf-8")
    calls=[]
    m=SimpleNamespace(
        DATA_ROOT=tmp_path,OPS_FILE=ops,
        load_ops=lambda:{"recovery_fifo":[]},
        _get_recovery_fifo_rows=lambda data:data.get("recovery_fifo",[]),
        _xlsx_stream_book=_xlsx_stream_book,
        _xlsx_monthly_rows=_xlsx_monthly_rows,
        _monthly_sheet_name=_monthly_sheet_name,
        _V368_start_internal_rebuild=lambda mode,value:(
            calls.append((mode,value)) or {"status":"queued","job_id":"safe-job"}
        ),
    )

    class RunImmediately:
        def __init__(self,target,**kwargs):self.target=target
        def start(self):self.target()

    with patch("v387_one_time_dev_repair.Thread",RunImmediately),patch(
            "v387_one_time_dev_repair.time.sleep",lambda secs:None):
        install(m)
        install(m)
    assert calls==[("day","2026-10-25")]
    marker=json.loads((tmp_path/MARKER_NAME).read_text(encoding="utf-8"))
    assert marker["status"]=="queued"
    assert marker["job_id"]=="safe-job"
    assert "fingerprint" in marker


def test_no_persisted_fifo_change_outside_atomic_v368_worker():
    root=Path(__file__).resolve().parents[1]
    code=(root/"v387_one_time_dev_repair.py").read_text(encoding="utf-8")
    assert "_V368_start_internal_rebuild" in code
    assert "os.replace(temp,path)" in code
    assert "m.OPS_FILE.write_" not in code
    worker=(root/"v368_rebuild_dev_patch.py").read_text(encoding="utf-8")
    assert "os.link(m.OPS_FILE, backup)" in worker
    assert "temporary.replace(m.OPS_FILE)" in worker
    assert "if new_dev <= previous_dev:" in worker
