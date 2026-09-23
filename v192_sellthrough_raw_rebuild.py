"""V192 · Sell Through reconstruido desde el Excel fuente.

Esta versión NO usa el dataframe/cache normalizado de Comercial para inventario.
Lee directamente el archivo XLSX de capacidades con openpyxl en modo streaming,
crea un cache SQLite pequeño y calcula Sell Through desde las columnas fuente.
"""
from __future__ import annotations

import json
import math
import sqlite3
import threading
import time
import unicodedata
from pathlib import Path

from fastapi import Request, HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def install(m):
    if getattr(m, "_V192_SELLTHROUGH_RAW", False):
        return

    def norm(v):
        t=unicodedata.normalize("NFKD",str(v or ""))
        return " ".join("".join(c for c in t if not unicodedata.combining(c)).casefold().split())

    def clean_id(v):
        if v is None:return ""
        if isinstance(v,(int,)):
            return str(v)
        if isinstance(v,float) and math.isfinite(v) and v.is_integer():
            return str(int(v))
        s=str(v).strip()
        return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s

    def num(v):
        if v is None:return 0.0
        try:
            if isinstance(v,str):
                s=v.strip().replace("$","").replace(" ","")
                if not s:return 0.0
                s=s.replace(",","")
                x=float(s)
            else:
                x=float(v)
            return x if math.isfinite(x) else 0.0
        except Exception:
            return 0.0

    def section_group(v):
        k=norm(v)
        if "dama" in k or "mujer" in k:return "Dama"
        if "caballero" in k or "hombre" in k:return "Caballero"
        if "infantil" in k or "nino" in k or "niño" in k:return "Infantil"
        return str(v or "Sin sección").strip() or "Sin sección"

    raw_db=m.DATA_ROOT/"sellthrough_raw_v2.sqlite3"
    raw_lock=threading.RLock()
    RAW_VERSION="2"

    sales_db=m.DATA_ROOT/"sellthrough_sales_by_id.sqlite3"

    def source_stamp(path:Path):
        try:
            st=path.stat()
            return f"{st.st_mtime_ns}:{st.st_size}"
        except Exception:
            return ""

    def db_meta(path:Path):
        if not path.exists():return {}
        try:
            con=sqlite3.connect(path)
            try:return {str(k):str(v) for k,v in con.execute("SELECT k,v FROM meta")}
            finally:con.close()
        except Exception:return {}

    def pick_index(headers, aliases, last=False):
        alias_keys={norm(x) for x in aliases}
        matches=[i for i,h in enumerate(headers) if norm(h) in alias_keys]
        if not matches:return None
        return matches[-1] if last else matches[0]

    def header_map(headers):
        # EXISTENCIA: elegir la última columna EXACTA. En el archivo actual es AO.
        exact_exist=[i for i,h in enumerate(headers) if norm(h)=="existencia"]
        idx={
            "id":pick_index(headers,["ID_ART","ID ART","ID","CODIGO"]),
            "store":pick_index(headers,["TIENDA","SUCURSAL","TIENDA/SUCURSAL"]),
            "model":pick_index(headers,["MODELO"]),
            "brand":pick_index(headers,["MARCA PRICE","MARCA"]),
            "section":pick_index(headers,["SECCION","SECCIÓN"]),
            "rubro":pick_index(headers,["SUBCATEGORIA","SUBCATEGORÍA","RUBRO"]),
            "existence":(exact_exist[-1] if exact_exist else pick_index(headers,["EXISTENCIA TOTAL"],last=True)),
            "cedis":pick_index(headers,["EXISTENCIA CEDIS","EXISTENCIA EN CEDIS","EXISTENCIA CEDIS PZAS","EXIST CEDIS","INVENTARIO CEDIS"]),
            "transit":pick_index(headers,["TRANSITO","TRÁNSITO","EN TRANSITO","EN TRÁNSITO","MERCANCIA EN TRANSITO","MERCANCÍA EN TRÁNSITO"]),
            "suggested":pick_index(headers,["SUG 7","SUGERIDO 7","VPD"]),
            "catalog":pick_index(headers,["TIPO CATALOGO MAX VIG","TIPO CATÁLOGO MAX VIG","TIPO CATALOGO","TIPO CATÁLOGO"]),
            "status":pick_index(headers,["ESTATUS DE CATALOGO","ESTATUS DE CATÁLOGO","ESTATUS CATALOGO","ESTATUS CATÁLOGO","ESTATUS CATALOGO MAX VIG","ESTATUS CATÁLOGO MAX VIG"]),
        }
        return idx

    def row_value(row,idx):
        if idx is None or idx<0 or idx>=len(row):return None
        return row[idx]

    def is_vigente(status_value,catalog_value):
        sk=norm(status_value)
        ck=norm(catalog_value)
        def good(k):
            return bool(k) and ("vigente" in k or k=="vig") and not any(x in k for x in ("no vigente","descontinu","inactiv","baja","cancel","suspend"))
        if sk:
            return 1 if good(sk) else 0
        if ck:
            return 1 if good(ck) else 0
        # Archivo histórico sin campo de vigencia: no excluir el modelo.
        return 1

    def raw_ready(entry,path):
        meta=db_meta(raw_db)
        return (
            meta.get("version")==RAW_VERSION
            and meta.get("source_id")==str(entry.get("id") or "")
            and meta.get("source_stamp")==source_stamp(path)
        )

    def build_raw_cache(entry,path):
        with raw_lock:
            if raw_ready(entry,path):return
            tmp=raw_db.with_suffix(".tmp.sqlite3")
            try:tmp.unlink(missing_ok=True)
            except Exception:pass
            con=sqlite3.connect(tmp)
            try:
                con.executescript("""
                    PRAGMA journal_mode=OFF;
                    PRAGMA synchronous=OFF;
                    PRAGMA temp_store=FILE;
                    CREATE TABLE meta(k TEXT PRIMARY KEY,v TEXT NOT NULL);
                    CREATE TABLE inventory(
                        store TEXT NOT NULL,
                        id_art TEXT NOT NULL,
                        model TEXT NOT NULL DEFAULT '',
                        brand TEXT NOT NULL DEFAULT '',
                        section TEXT NOT NULL DEFAULT '',
                        rubro TEXT NOT NULL DEFAULT '',
                        existence REAL NOT NULL DEFAULT 0,
                        cedis REAL NOT NULL DEFAULT 0,
                        transit REAL NOT NULL DEFAULT 0,
                        suggested REAL NOT NULL DEFAULT 0,
                        vigente INTEGER NOT NULL DEFAULT 0,
                        PRIMARY KEY(store,id_art)
                    );
                    CREATE INDEX idx_raw_id ON inventory(id_art);
                    CREATE INDEX idx_raw_store ON inventory(store);
                    CREATE INDEX idx_raw_vig ON inventory(vigente);
                """)
                wb=load_workbook(path,read_only=True,data_only=True)
                total_rows=0
                parsed_sheets=0
                header_notes=[]
                try:
                    for ws in wb.worksheets:
                        rows=ws.iter_rows(values_only=True)
                        headers=None;idx=None;header_row=0
                        for n,row in enumerate(rows,start=1):
                            vals=list(row or [])
                            candidate=[str(x or "").strip() for x in vals]
                            keys={norm(x) for x in candidate if str(x or "").strip()}
                            if ({"id_art","id art"} & keys) and ("existencia" in keys or "existencia total" in keys):
                                headers=candidate;idx=header_map(headers);header_row=n;break
                            if n>=80:break
                        if not headers or idx.get("id") is None or idx.get("existence") is None:
                            continue
                        parsed_sheets+=1
                        note={
                            "sheet":ws.title,
                            "header_row":header_row,
                            "id":f"{get_column_letter(idx['id']+1)}:{headers[idx['id']]}",
                            "store":f"{get_column_letter(idx['store']+1)}:{headers[idx['store']]}" if idx.get("store") is not None else "",
                            "existence":f"{get_column_letter(idx['existence']+1)}:{headers[idx['existence']]}",
                            "transit":f"{get_column_letter(idx['transit']+1)}:{headers[idx['transit']]}" if idx.get("transit") is not None else "",
                            "status":f"{get_column_letter(idx['status']+1)}:{headers[idx['status']]}" if idx.get("status") is not None else "",
                            "catalog":f"{get_column_letter(idx['catalog']+1)}:{headers[idx['catalog']]}" if idx.get("catalog") is not None else "",
                        }
                        header_notes.append(note)
                        batch=[]
                        for row in rows:
                            vals=list(row or [])
                            ident=clean_id(row_value(vals,idx["id"]))
                            if not ident or ident.lower() in ("nan","none"):continue
                            store=m._normalize_store_value(row_value(vals,idx.get("store"))) if idx.get("store") is not None else ""
                            model=str(row_value(vals,idx.get("model")) or ident).strip()
                            brand=str(row_value(vals,idx.get("brand")) or "").strip()
                            sec=section_group(row_value(vals,idx.get("section")))
                            rubro=str(row_value(vals,idx.get("rubro")) or "").strip()
                            existence=num(row_value(vals,idx.get("existence")))
                            cedis=num(row_value(vals,idx.get("cedis")))
                            transit=num(row_value(vals,idx.get("transit")))
                            suggested=num(row_value(vals,idx.get("suggested")))
                            vigente=is_vigente(row_value(vals,idx.get("status")),row_value(vals,idx.get("catalog")))
                            batch.append((store,ident,model,brand,sec,rubro,existence,cedis,transit,suggested,vigente))
                            total_rows+=1
                            if len(batch)>=5000:
                                con.executemany("""
                                  INSERT INTO inventory(store,id_art,model,brand,section,rubro,existence,cedis,transit,suggested,vigente)
                                  VALUES(?,?,?,?,?,?,?,?,?,?,?)
                                  ON CONFLICT(store,id_art) DO UPDATE SET
                                    model=CASE WHEN inventory.model='' THEN excluded.model ELSE inventory.model END,
                                    brand=CASE WHEN inventory.brand='' THEN excluded.brand ELSE inventory.brand END,
                                    section=CASE WHEN inventory.section='' THEN excluded.section ELSE inventory.section END,
                                    rubro=CASE WHEN inventory.rubro='' THEN excluded.rubro ELSE inventory.rubro END,
                                    existence=inventory.existence+excluded.existence,
                                    cedis=MAX(inventory.cedis,excluded.cedis),
                                    transit=inventory.transit+excluded.transit,
                                    suggested=inventory.suggested+excluded.suggested,
                                    vigente=MAX(inventory.vigente,excluded.vigente)
                                """,batch)
                                con.commit();batch.clear()
                        if batch:
                            con.executemany("""
                              INSERT INTO inventory(store,id_art,model,brand,section,rubro,existence,cedis,transit,suggested,vigente)
                              VALUES(?,?,?,?,?,?,?,?,?,?,?)
                              ON CONFLICT(store,id_art) DO UPDATE SET
                                model=CASE WHEN inventory.model='' THEN excluded.model ELSE inventory.model END,
                                brand=CASE WHEN inventory.brand='' THEN excluded.brand ELSE inventory.brand END,
                                section=CASE WHEN inventory.section='' THEN excluded.section ELSE inventory.section END,
                                rubro=CASE WHEN inventory.rubro='' THEN excluded.rubro ELSE inventory.rubro END,
                                existence=inventory.existence+excluded.existence,
                                cedis=MAX(inventory.cedis,excluded.cedis),
                                transit=inventory.transit+excluded.transit,
                                suggested=inventory.suggested+excluded.suggested,
                                vigente=MAX(inventory.vigente,excluded.vigente)
                            """,batch)
                            con.commit()
                finally:
                    wb.close()

                meta=[
                    ("version",RAW_VERSION),
                    ("source_id",str(entry.get("id") or "")),
                    ("source_stamp",source_stamp(path)),
                    ("source_file",path.name),
                    ("rows",str(total_rows)),
                    ("sheets",str(parsed_sheets)),
                    ("headers",json.dumps(header_notes,ensure_ascii=False)),
                    ("built_at",str(time.time())),
                ]
                con.executemany("INSERT OR REPLACE INTO meta(k,v) VALUES(?,?)",meta)
                con.commit()

                dbg=con.execute("SELECT COUNT(*),COALESCE(SUM(existence),0),COALESCE(SUM(transit),0),COALESCE(MAX(cedis),0) FROM inventory WHERE id_art='1322682'").fetchone()
                print(f"[V192-RAW-SELFTEST] id=1322682 tiendas={dbg[0]} existencia={dbg[1]:.0f} transito={dbg[2]:.0f} cedis={dbg[3]:.0f}",flush=True)
                print(f"[V192-RAW-HEADERS] {json.dumps(header_notes,ensure_ascii=False)}",flush=True)
            finally:
                con.close()
            tmp.replace(raw_db)
            print(f"[V192-RAW] Cache directo creado desde {path.name} filas={total_rows}",flush=True)

    def ensure_raw(week):
        entry=m._capacity_source_entry(week)
        if not entry:raise HTTPException(503,"No hay Excel de capacidades procesado")
        path=m.resolve_entry_path(entry)
        if not path.exists():raise HTTPException(503,"No se encontró el Excel fuente de capacidades")
        if not raw_ready(entry,path):build_raw_cache(entry,path)
        return entry,path

    def sales_map(selected_store,company_scope):
        if not sales_db.exists():
            raise HTTPException(503,"Las ventas acumuladas aún no están preparadas")
        con=sqlite3.connect(sales_db)
        try:
            meta={str(k):str(v) for k,v in con.execute("SELECT k,v FROM meta")}
            if company_scope:
                rows=con.execute("SELECT id_art,SUM(sales) FROM sales GROUP BY id_art").fetchall()
            else:
                rows=con.execute("SELECT id_art,SUM(sales) FROM sales WHERE store=? GROUP BY id_art",(selected_store,)).fetchall()
        finally:con.close()
        try:months=json.loads(meta.get("months") or "[]")
        except Exception:months=[]
        return {str(a):float(v or 0) for a,v in rows},months

    def inventory_rows(selected_store,company_scope):
        con=sqlite3.connect(raw_db)
        con.row_factory=sqlite3.Row
        try:
            eligible={str(r[0]) for r in con.execute("SELECT id_art FROM inventory GROUP BY id_art HAVING MAX(vigente)=1")}
            if not eligible:return []
            if company_scope:
                rows=con.execute("""
                    SELECT id_art,
                           MIN(CASE WHEN model<>'' THEN model END) model,
                           MIN(CASE WHEN brand<>'' THEN brand END) brand,
                           MIN(CASE WHEN section<>'' THEN section END) section,
                           MIN(CASE WHEN rubro<>'' THEN rubro END) rubro,
                           SUM(existence) existence,
                           MAX(cedis) cedis,
                           SUM(transit) transit,
                           SUM(suggested) suggested
                    FROM inventory
                    GROUP BY id_art
                """).fetchall()
            else:
                rows=con.execute("""
                    SELECT id_art,model,brand,section,rubro,existence,cedis,transit,suggested
                    FROM inventory WHERE store=?
                """,(selected_store,)).fetchall()
            return [dict(r) for r in rows if str(r["id_art"]) in eligible]
        finally:con.close()

    @m.app.get("/api/commercial-sellthrough-v2")
    def sellthrough_v2(request:Request,week:str="",store:str="Compañía",section:str="Todas",catalog:str="Todos"):
        actor=m.require_user(request)
        selected_store=str(m.effective_store(actor,store) or "Compañía")
        company_scope=norm(selected_store)==norm("Compañía")
        entry,path=ensure_raw(week)
        sales,months=sales_map(selected_store,company_scope)
        inv=inventory_rows(selected_store,company_scope)

        rows=[]
        for r in inv:
            sec=str(r.get("section") or "Sin sección")
            if section!="Todas" and norm(sec)!=norm(section):continue
            ident=str(r.get("id_art") or "")
            sale=float(sales.get(ident,0.0) or 0.0)
            existence=float(r.get("existence") or 0.0)
            cedis=float(r.get("cedis") or 0.0)
            transit=float(r.get("transit") or 0.0)
            stock=existence+transit+(cedis if company_scope else 0.0)
            base=sale+stock
            st=(sale/base*100.0) if base>0 else 0.0
            rows.append({
                "id_art":ident,"model":str(r.get("model") or ident),"brand":str(r.get("brand") or ""),
                "section":sec,"rubro":str(r.get("rubro") or ""),"sales_pzas":sale,
                "existence":existence,"cedis_existence":cedis,"transit":transit,
                "stock_available":stock,"inventory_for_sell":stock,"available_base":base,
                "suggested":float(r.get("suggested") or 0.0),"sell_through":st,
            })
        rows.sort(key=lambda x:(-x["sell_through"],-x["sales_pzas"],x["id_art"]))

        total_sales=sum(x["sales_pzas"] for x in rows)
        total_exist=sum(x["existence"] for x in rows)
        total_trans=sum(x["transit"] for x in rows)
        # CEDIS ya viene una vez por ID en rows; por tanto sí se suma entre modelos.
        total_cedis=sum(x["cedis_existence"] for x in rows)
        total_stock=total_exist+total_trans+(total_cedis if company_scope else 0.0)
        total_base=total_sales+total_stock
        total_st=total_sales/total_base*100 if total_base>0 else 0.0
        totals={
            "sell_through":total_st,"sales_pzas":total_sales,"existence":total_exist,
            "cedis_existence":total_cedis,"transit":total_trans,"stock_available":total_stock,
            "inventory_for_sell":total_stock,"available_base":total_base,
            "models":len(rows),"models_ge_80":sum(1 for x in rows if x["sell_through"]>=80),
        }
        if any(x["id_art"]=="1322682" for x in rows):
            x=next(x for x in rows if x["id_art"]=="1322682")
            print(f"[V192-ENDPOINT-ID] id=1322682 existencia={x['existence']:.0f} venta={x['sales_pzas']:.0f} stock={x['stock_available']:.0f} st={x['sell_through']:.2f}",flush=True)

        return {
            "week":week,"store":selected_store,"section":section,"catalog":"TIPO CATALOGO MAX VIG",
            "status":"VIGENTE","rows":rows,"totals":totals,
            "formula":"Vta acum pzs / (Vta acum pzs + Stock disponible)",
            "source":"Excel de capacidades fuente + Base de muertos y cambios",
            "source_file":path.name,"sales_scope":"Acumulado "+(" + ".join(months) if months else "meses disponibles"),
            "cedis_in_sellthrough":company_scope,"engine":"RAW-V2",
        }

    # Precalentar el cache directo al iniciar para validar el archivo antes de que
    # el usuario abra la pestaña.
    def warm():
        try:
            time.sleep(8)
            entry=m._capacity_source_entry("")
            if entry:
                path=m.resolve_entry_path(entry)
                if path.exists():build_raw_cache(entry,path)
        except Exception as exc:
            print(f"[V192-RAW] Error precalentando: {type(exc).__name__}: {exc}",flush=True)
    threading.Thread(target=warm,name="sellthrough-raw-v2-warm",daemon=True).start()

    m._V192_SELLTHROUGH_RAW=True
    print("[V192] Sell Through RAW-V2 instalado: inventario leído directo del XLSX.",flush=True)
